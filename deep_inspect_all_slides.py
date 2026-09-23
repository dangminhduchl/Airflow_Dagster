import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def hex_color(rgb_color):
    if rgb_color is None:
        return "None"
    try:
        return f"#{rgb_color[0]:02X}{rgb_color[1]:02X}{rgb_color[2]:02X}"
    except Exception:
        return str(rgb_color)

def parse_shape(shp, depth=0):
    indent = "  " * depth
    info = {
        "name": shp.name,
        "type": str(shp.shape_type),
        "fill": None,
        "line": None,
        "line_w": None,
        "radius": None,
        "text": None,
        "x": round(shp.left.inches, 2) if hasattr(shp, 'left') else None,
        "y": round(shp.top.inches, 2) if hasattr(shp, 'top') else None,
        "w": round(shp.width.inches, 2) if hasattr(shp, 'width') else None,
        "h": round(shp.height.inches, 2) if hasattr(shp, 'height') else None,
    }

    if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = []
        for child in shp.shapes:
            children.extend(parse_shape(child, depth + 1))
        return children

    # Text
    if shp.has_text_frame:
        txt = " ".join([p.text.strip() for p in shp.text_frame.paragraphs if p.text.strip()])
        if txt:
            info["text"] = txt[:40]

    # Fill
    try:
        if shp.fill.type == 1: # solid
            info["fill"] = hex_color(shp.fill.fore_color.rgb)
        elif shp.fill.type == 0:
            info["fill"] = "NoFill"
        elif shp.fill.type is not None:
            info["fill"] = f"Type_{shp.fill.type}"
    except Exception:
        pass

    # Line
    try:
        if shp.line.fill.type == 1:
            info["line"] = hex_color(shp.line.color.rgb)
            if shp.line.width:
                info["line_w"] = f"{shp.line.width.pt:.1f}pt"
        elif shp.line.fill.type == 0:
            info["line"] = "NoLine"
    except Exception:
        pass

    # Radius
    try:
        if len(shp.adjustments) > 0:
            info["radius"] = round(float(shp.adjustments[0]), 4)
    except Exception:
        pass

    # Also inspect XML directly for background and borders if python-pptx doesn't expose it
    spPr = shp.element.spPr
    if spPr is not None:
        # Check solidFill in XML
        solidFill = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        if solidFill is not None:
            srgbClr = solidFill.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
            if srgbClr is not None and not info["fill"]:
                info["fill"] = f"#{srgbClr.get('val')}"
        # Check ln in XML
        ln = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        if ln is not None:
            if not info["line_w"] and ln.get('w'):
                info["line_w"] = f"{int(ln.get('w'))/12700:.1f}pt"
            ln_solid = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
            if ln_solid is not None:
                ln_srgb = ln_solid.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
                if ln_srgb is not None and not info["line"]:
                    info["line"] = f"#{ln_srgb.get('val')}"

    return [info]

prs = Presentation('Airflow_vs_Dagster_Thuyet_Trinh.pptx.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1} ====================")
    # Check slide background in XML
    cSld = slide.element.cSld
    bg = cSld.bg
    bg_color = "Default/None"
    if bg is not None:
        bgPr = bg.bgPr
        if bgPr is not None:
            solidFill = bgPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
            if solidFill is not None:
                srgb = solidFill.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
                if srgb is not None:
                    bg_color = f"#{srgb.get('val')}"
    print(f"Slide BG Color: {bg_color}")
    
    all_elements = []
    for shp in slide.shapes:
        all_elements.extend(parse_shape(shp))
    
    # Analyze cards / shapes
    fills = set()
    lines = set()
    radiuses = set()
    for el in all_elements:
        if el["fill"]: fills.add(el["fill"])
        if el["line"]: lines.add(f"{el['line']} ({el['line_w']})")
        if el["radius"] is not None: radiuses.add(el["radius"])
    
    print(f"Unique Fills: {sorted(list(fills))}")
    print(f"Unique Lines: {sorted(list(lines))}")
    print(f"Unique Radiuses: {sorted(list(radiuses))}")
    
    # List top-level or significant cards
    print("Detailed Elements with Line or Fill:")
    for el in all_elements:
        if el["fill"] not in [None, "NoFill"] or el["line"] not in [None, "NoLine"]:
            print(f"  - {el['name']} ({el['type']}) x={el['x']}\" y={el['y']}\" w={el['w']}\" h={el['h']}\"")
            print(f"      fill={el['fill']}, line={el['line']} ({el['line_w']}), radius={el['radius']}, text={el['text']}")
