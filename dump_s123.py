from pptx import Presentation

prs = Presentation("Airflow_vs_Dagster_Thuyet_Trinh.pptx.pptx")
print(f"File: Airflow_vs_Dagster_Thuyet_Trinh.pptx.pptx, Size: {prs.slide_width.inches}x{prs.slide_height.inches}")

for idx in [0, 1, 2]:
    slide = prs.slides[idx]
    print(f"\n--- Slide {idx+1} ({len(slide.shapes)} shapes) ---")
    for s in slide.shapes:
        t = s.text_frame.text.replace('\n', ' ')[:60] if s.has_text_frame else "[Non-text]"
        print(f"  Shape {s.name} ({s.shape_type}): pos=({s.left.inches:.2f}, {s.top.inches:.2f}), size=({s.width.inches:.2f}x{s.height.inches:.2f}) -> {t}")
