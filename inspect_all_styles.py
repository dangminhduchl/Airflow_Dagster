import sys
import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE

def hex_color(rgb_color):
    if rgb_color is None:
        return "None"
    try:
        return f"#{rgb_color[0]:02X}{rgb_color[1]:02X}{rgb_color[2]:02X}"
    except Exception:
        return str(rgb_color)

def inspect_presentation(path):
    prs = Presentation(path)
    print(f"=== PRESENTATION: {path} ===")
    print(f"Dimensions: {prs.slide_width.inches:.2f}\" x {prs.slide_height.inches:.2f}\"")
    print(f"Total Slides: {len(prs.slides)}\n")

    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        print(f"--- SLIDE {slide_num} ---")
        
        # Check background
        bg = slide.background
        bg_fill = "default"
        if bg and bg.fill:
            try:
                bg_fill = f"type={bg.fill.type}, color={hex_color(bg.fill.fore_color.rgb)}"
            except Exception as e:
                bg_fill = f"error/theme ({e})"
        print(f"Slide Background: {bg_fill}")
        
        print(f"Total shapes: {len(slide.shapes)}")
        
        shape_stats = {
            "rounded_rectangles": [],
            "rectangles": [],
            "textboxes": [],
            "pictures": [],
            "others": []
        }

        for s_idx, shp in enumerate(slide.shapes):
            shp_info = {
                "idx": s_idx,
                "name": shp.name,
                "type": str(shp.shape_type),
                "x": round(shp.left.inches, 2) if hasattr(shp, 'left') else 0,
                "y": round(shp.top.inches, 2) if hasattr(shp, 'top') else 0,
                "w": round(shp.width.inches, 2) if hasattr(shp, 'width') else 0,
                "h": round(shp.height.inches, 2) if hasattr(shp, 'height') else 0,
                "fill": "None",
                "line_color": "None",
                "line_width": "None",
                "radius": "N/A",
                "text": ""
            }

            if shp.has_text_frame:
                txt = " | ".join([p.text.strip() for p in shp.text_frame.paragraphs if p.text.strip()])
                shp_info["text"] = txt[:60]

            # Fill
            try:
                if shp.fill.type == 1: # solid
                    shp_info["fill"] = hex_color(shp.fill.fore_color.rgb)
                elif shp.fill.type == 0:
                    shp_info["fill"] = "Background/NoFill"
                else:
                    shp_info["fill"] = f"Type_{shp.fill.type}"
            except Exception:
                shp_info["fill"] = "None/Inherited"

            # Line
            try:
                if shp.line.fill.type == 1:
                    shp_info["line_color"] = hex_color(shp.line.color.rgb)
                    if shp.line.width:
                        shp_info["line_width"] = f"{shp.line.width.pt:.1f}pt"
                elif shp.line.fill.type == 0:
                    shp_info["line_color"] = "NoLine"
                else:
                    shp_info["line_color"] = "Default"
            except Exception:
                shp_info["line_color"] = "None"

            # Radius (adjustments)
            if shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    if len(shp.adjustments) > 0:
                        shp_info["radius"] = round(float(shp.adjustments[0]), 4)
                except Exception:
                    pass

            # Classify
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                shape_stats["pictures"].append(shp_info)
            elif shp.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                shape_stats["textboxes"].append(shp_info)
            elif shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                if "Round" in shp.name or shp_info["radius"] != "N/A":
                    shape_stats["rounded_rectangles"].append(shp_info)
                else:
                    shape_stats["rectangles"].append(shp_info)
            else:
                shape_stats["others"].append(shp_info)

        print(f"  Summary: {len(shape_stats['rounded_rectangles'])} Rounded Rects, {len(shape_stats['rectangles'])} Rects, {len(shape_stats['textboxes'])} Textboxes, {len(shape_stats['pictures'])} Pictures, {len(shape_stats['others'])} Others")
        
        # Detail on shapes with fills or lines or radiuses
        all_cards = shape_stats["rounded_rectangles"] + shape_stats["rectangles"]
        for card in all_cards:
            print(f"    - [{card['idx']}] {card['name']} (x={card['x']}\", y={card['y']}\", w={card['w']}\", h={card['h']}\")")
            print(f"        Fill: {card['fill']} | Line: {card['line_color']} ({card['line_width']}) | Radius: {card['radius']} | Text: {card['text']}")
        for pic in shape_stats["pictures"]:
            print(f"    - [PIC {pic['idx']}] (x={pic['x']}\", y={pic['y']}\", w={pic['w']}\", h={pic['h']}\")")
        print()

if __name__ == "__main__":
    inspect_presentation("Airflow_vs_Dagster_Thuyet_Trinh.pptx.pptx")
