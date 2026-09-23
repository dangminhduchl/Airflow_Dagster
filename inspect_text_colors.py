import pptx
from pptx import Presentation

def hex_color(rgb_color):
    if rgb_color is None:
        return "None"
    try:
        return f"#{rgb_color[0]:02X}{rgb_color[1]:02X}{rgb_color[2]:02X}"
    except Exception:
        return str(rgb_color)

def inspect_text_colors(path):
    prs = Presentation(path)
    for idx, slide in enumerate(prs.slides):
        print(f"\n==================== SLIDE {idx+1} TEXT COLORS ====================")
        text_color_counts = {}
        text_samples = {}
        
        # Traverse shapes recursively
        def parse_texts(shp):
            if shp.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                for c in shp.shapes:
                    parse_texts(c)
                return

            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    for r in p.runs:
                        clr = "Inherited/Default"
                        try:
                            if r.font.color and r.font.color.rgb:
                                clr = hex_color(r.font.color.rgb)
                        except Exception:
                            pass
                        
                        # Also check rPr in XML
                        rPr = r._r.rPr
                        if rPr is not None and clr == "Inherited/Default":
                            solidFill = rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
                            if solidFill is not None:
                                srgb = solidFill.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
                                if srgb is not None:
                                    clr = f"#{srgb.get('val').upper()}"

                        text_color_counts[clr] = text_color_counts.get(clr, 0) + 1
                        if clr not in text_samples:
                            text_samples[clr] = []
                        if len(text_samples[clr]) < 3 and r.text.strip():
                            text_samples[clr].append(r.text.strip()[:35])

        for shp in slide.shapes:
            parse_texts(shp)

        print(f"Colors used on Slide {idx+1}:")
        for clr, count in sorted(text_color_counts.items(), key=lambda x: -x[1]):
            samples = " | ".join(text_samples.get(clr, []))
            print(f"   {clr:20s}: {count:3d} runs -> e.g. '{samples}'")

if __name__ == "__main__":
    inspect_text_colors("Airflow_vs_Dagster_Thuyet_Trinh.pptx.pptx")
