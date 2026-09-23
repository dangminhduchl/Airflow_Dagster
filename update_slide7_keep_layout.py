import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def build_slide7():
    pptx_path = "Airflow_vs_Dagster_Thuyet_Trinh.pptx.pptx"
    prs = Presentation(pptx_path)
    
    # 20.0" x 11.25"
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    
    # Palette matching Slide 1, 2, 3
    BG_COLOR = RGBColor(7, 5, 14)          # #07050e
    CARD_BG = RGBColor(18, 14, 32)         # #120e20
    CARD_BG_ALT = RGBColor(15, 23, 42)     # #0f172a
    CARD_BG_CODE = RGBColor(5, 8, 20)      # #050814
    
    ACCENT_CYAN = RGBColor(56, 189, 248)   # #38bdf8
    ACCENT_PURPLE = RGBColor(192, 132, 252)# #c084fc
    ACCENT_GREEN = RGBColor(34, 197, 94)   # #22c55e
    ACCENT_ROSE = RGBColor(244, 63, 94)    # #f43f5e
    ACCENT_AMBER = RGBColor(245, 158, 11)  # #f59e0b
    
    BORDER_PURPLE = RGBColor(88, 28, 135)  # #581c87
    BORDER_CYAN = RGBColor(14, 116, 144)   # #0e7490
    BORDER_SUBTLE = RGBColor(39, 31, 61)   # #271f3d
    BORDER_CODE = RGBColor(51, 65, 85)     # #334155
    
    TEXT_MAIN = RGBColor(248, 250, 252)    # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94a3b8
    FONT_MAIN = "Arial"
    FONT_CODE = "Consolas"

    s7 = prs.slides[6]
    
    # Clear all shapes
    for shape in list(s7.shapes):
        sp = shape._element
        sp.getparent().remove(sp)
        
    # Helper: Card creation
    def add_card(x, y, w, h, border_color=BORDER_PURPLE, bg_color=CARD_BG, border_width=1.5, radius=0.05):
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
        card.adjustments[0] = radius
        return card

    # Helper: Pill badge
    def add_badge(x, y, w, h, text, text_color, border_color, bg_color=CARD_BG_ALT, font_size=10.0):
        b = add_card(x, y, w, h, border_color=border_color, bg_color=bg_color, border_width=1.0, radius=0.5)
        tf = b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = FONT_MAIN
        p.font.size = Pt(font_size)
        p.font.bold = True
        p.font.color.rgb = text_color
        p.alignment = PP_ALIGN.CENTER
        return b

    # Slide Background
    bg = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_width, slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()

    # 1. Slide Header
    add_badge(Inches(0.98), Inches(0.38), Inches(3.6), Inches(0.48),
              "TRIỂN KHAI & VẬN HÀNH", ACCENT_PURPLE, BORDER_PURPLE, CARD_BG, 11.5)

    title_box = s7.shapes.add_textbox(Inches(4.88), Inches(0.32), Inches(14.2), Inches(0.55))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = Inches(0)
    p_t = tf_t.paragraphs[0]
    p_t.text = "HẠ TẦNG ON-PREMISE & ĐỊNH VỊ NỀN TẢNG"
    p_t.font.name = FONT_MAIN
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_MAIN

    p_sub = tf_t.add_paragraph()
    p_sub.text = "Kiến trúc cụm On-Premise, năng lực kiểm thử CI/CD và định vị lựa chọn công nghệ"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = ACCENT_CYAN
    p_sub.space_before = Pt(4)

    # Coordinates
    x_left = Inches(0.98)
    col_w = Inches(8.84)
    gap_col = Inches(0.36)
    x_right = x_left + col_w + gap_col  # 10.18"
    
    top_y = Inches(1.15)
    top_h = Inches(4.70)
    
    bot_y = Inches(6.05)
    bot_w = Inches(18.04)
    bot_h = Inches(4.80)

    # ================= BLOCK 1 (TOP LEFT): INFRASTRUCTURE =================
    b1_card = add_card(x_left, top_y, col_w, top_h, border_color=ACCENT_CYAN, bg_color=CARD_BG, border_width=1.5, radius=0.05)
    
    # Title & Badge
    b1_tbox = s7.shapes.add_textbox(x_left + Inches(0.30), top_y + Inches(0.18), Inches(6.0), Inches(0.35))
    tf = b1_tbox.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "🏗️ KIẾN TRÚC HẠ TẦNG ON-PREMISE"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    add_badge(x_left + col_w - Inches(1.80), top_y + Inches(0.18), Inches(1.50), Inches(0.32),
              "Ops Overhead", ACCENT_CYAN, BORDER_CYAN, CARD_BG_ALT, 10.0)

    # Sub-card 1: Airflow Infra
    s1_w = col_w - Inches(0.56)
    s1_x = x_left + Inches(0.28)
    s1_y = top_y + Inches(0.65)
    s1_h = Inches(1.85)
    add_card(s1_x, s1_y, s1_w, s1_h, border_color=ACCENT_ROSE, bg_color=CARD_BG_ALT, border_width=1.0, radius=0.04)
    
    s1_tbox = s7.shapes.add_textbox(s1_x + Inches(0.20), s1_y + Inches(0.12), s1_w - Inches(0.40), s1_h - Inches(0.24))
    tf = s1_tbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "🔴 AIRFLOW (6 THÀNH PHẦN RỜI RẠC)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    pr = p.add_run()
    pr.text = "   — Cồng kềnh & Dễ nghẽn"
    pr.font.bold = False
    pr.font.size = Pt(9.5)
    pr.font.color.rgb = RGBColor(252, 165, 165)

    bullets_s1 = [
        ("Stack phức tạp:", "Scheduler + Celery Worker Pool + Redis Queue + Postgres HA + Webserver + Triggerer."),
        ("Nỗi đau đĩa chia sẻ:", "Bắt buộc dựng NFS/EFS chia sẻ để đồng bộ file dags/ giữa các node, nguy cơ nghẽn IOPS."),
        ("Xung đột package:", "Toàn bộ task chung Python env; 1 team upgrade pandas làm gãy pipeline team khác.")
    ]
    for b_title, b_desc in bullets_s1:
        p = tf.add_paragraph()
        p.text = f"• {b_title} "
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(3)
        pr = p.add_run()
        pr.text = b_desc
        pr.font.bold = False
        pr.font.color.rgb = TEXT_MUTED

    # Sub-card 2: Dagster Infra
    s2_y = s1_y + s1_h + Inches(0.15)
    s2_h = Inches(1.85)
    add_card(s1_x, s2_y, s1_w, s2_h, border_color=ACCENT_GREEN, bg_color=CARD_BG_ALT, border_width=1.0, radius=0.04)

    s2_tbox = s7.shapes.add_textbox(s1_x + Inches(0.20), s2_y + Inches(0.12), s1_w - Inches(0.40), s2_h - Inches(0.24))
    tf = s2_tbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "🟢 DAGSTER (CHUẨN KUBERNETES & CLOUD-NATIVE)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    pr = p.add_run()
    pr.text = "   — Tinh gọn & Độc lập"
    pr.font.bold = False
    pr.font.size = Pt(9.5)
    pr.font.color.rgb = RGBColor(134, 239, 172)

    bullets_s2 = [
        ("Stack tinh gọn 3 dịch vụ:", "Chỉ cần Dagster Webserver + Daemon + PostgreSQL HA (chỉ lưu event metadata nhẹ)."),
        ("Lưu trữ đối tượng MinIO:", "Cụm MinIO phân tán làm S3 nội bộ, I/O Manager tự động đẩy/kéo dữ liệu."),
        ("Code Location Pods:", "Đóng gói Docker riêng, gọi gRPC. Tự do dùng Python 3.10/3.11/3.12 cách ly.")
    ]
    for b_title, b_desc in bullets_s2:
        p = tf.add_paragraph()
        p.text = f"• {b_title} "
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(3)
        pr = p.add_run()
        pr.text = b_desc
        pr.font.bold = False
        pr.font.color.rgb = TEXT_MUTED

    # ================= BLOCK 2 (TOP RIGHT): CI/CD & TESTING =================
    b2_card = add_card(x_right, top_y, col_w, top_h, border_color=ACCENT_PURPLE, bg_color=CARD_BG, border_width=1.5, radius=0.05)
    
    # Title & Badge
    b2_tbox = s7.shapes.add_textbox(x_right + Inches(0.30), top_y + Inches(0.18), Inches(6.0), Inches(0.35))
    tf = b2_tbox.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "🧪 NĂNG LỰC KIỂM THỬ CI/CD & DEVX"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    add_badge(x_right + col_w - Inches(2.00), top_y + Inches(0.18), Inches(1.70), Inches(0.32),
              "Pytest In-Memory", ACCENT_PURPLE, BORDER_PURPLE, CARD_BG_ALT, 10.0)

    # Sub-card R1: Airflow Testing
    r1_y = top_y + Inches(0.65)
    r1_h = Inches(1.50)
    add_card(x_right + Inches(0.28), r1_y, s1_w, r1_h, border_color=ACCENT_AMBER, bg_color=CARD_BG_ALT, border_width=1.0, radius=0.04)

    r1_tbox = s7.shapes.add_textbox(x_right + Inches(0.48), r1_y + Inches(0.12), s1_w - Inches(0.40), r1_h - Inches(0.24))
    tf = r1_tbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "⚠️ Airflow: Rào Cản Lớn Khi Viết Unit Test & CI/CD"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    bullets_r1 = [
        ("Phụ thuộc Context & DB:", "Task gắn chặt với ti, context, dag_run → Rất khó mock dữ liệu cục bộ."),
        ("Chu kỳ phản hồi chậm (3-5 phút):", "Muốn test phải dựng Docker Compose & migrate Postgres → Không thể tích hợp Git Hook.")
    ]
    for b_title, b_desc in bullets_r1:
        p = tf.add_paragraph()
        p.text = f"• {b_title} "
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(3)
        pr = p.add_run()
        pr.text = b_desc
        pr.font.bold = False
        pr.font.color.rgb = TEXT_MUTED

    # Sub-card R2: Dagster Testing + Code Block
    r2_y = r1_y + r1_h + Inches(0.15)
    r2_h = Inches(2.20)
    add_card(x_right + Inches(0.28), r2_y, s1_w, r2_h, border_color=BORDER_CYAN, bg_color=CARD_BG_ALT, border_width=1.0, radius=0.04)

    r2_tbox = s7.shapes.add_textbox(x_right + Inches(0.48), r2_y + Inches(0.10), s1_w - Inches(0.40), Inches(0.35))
    tf = r2_tbox.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "💎 Dagster: Unit Test Trong RAM 0.12s (4/4 Tests)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    add_badge(x_right + Inches(0.28) + s1_w - Inches(1.50), r2_y + Inches(0.08), Inches(1.30), Inches(0.26),
              "pytest passed", ACCENT_GREEN, RGBColor(22, 101, 52), CARD_BG_ALT, 9.0)

    # Code block inside R2
    code_w = s1_w - Inches(0.40)
    code_x = x_right + Inches(0.48)
    code_y = r2_y + Inches(0.45)
    code_h = Inches(1.15)
    add_card(code_x, code_y, code_w, code_h, border_color=BORDER_CODE, bg_color=CARD_BG_CODE, border_width=1.0, radius=0.03)

    code_tbox = s7.shapes.add_textbox(code_x + Inches(0.15), code_y + Inches(0.10), code_w - Inches(0.30), code_h - Inches(0.20))
    tf_c = code_tbox.text_frame
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = Inches(0)
    
    p = tf_c.paragraphs[0]
    p.font.name = FONT_CODE
    p.font.size = Pt(9.0)
    r1 = p.add_run(); r1.text = "def "; r1.font.color.rgb = ACCENT_PURPLE; r1.font.bold = True
    r2 = p.add_run(); r2.text = "test_budget_limit_blocks_over_budget():\n"; r2.font.color.rgb = RGBColor(96, 165, 250)
    
    p2 = tf_c.add_paragraph()
    p2.font.name = FONT_CODE
    p2.font.size = Pt(9.0)
    p2.space_before = Pt(2)
    r3 = p2.add_run(); r3.text = "  res = check_budget_limit_compliance(mock_120m_invoices)\n"; r3.font.color.rgb = TEXT_MAIN
    
    p3 = tf_c.add_paragraph()
    p3.font.name = FONT_CODE
    p3.font.size = Pt(9.0)
    p3.space_before = Pt(2)
    r4 = p3.add_run(); r4.text = "  assert not "; r4.font.color.rgb = ACCENT_ROSE; r4.font.bold = True
    r5 = p3.add_run(); r5.text = "res.passed  "; r5.font.color.rgb = TEXT_MAIN
    r6 = p3.add_run(); r6.text = "# Ngắt luồng trong 0.02s!"; r6.font.color.rgb = RGBColor(100, 116, 139)

    # Note footer
    foot_tbox = s7.shapes.add_textbox(code_x, code_y + code_h + Inches(0.06), code_w, Inches(0.40))
    tf_f = foot_tbox.text_frame
    tf_f.word_wrap = True
    tf_f.margin_left = tf_f.margin_right = tf_f.margin_top = tf_f.margin_bottom = Inches(0)
    pf = tf_f.paragraphs[0]
    pf.text = "✨ Tự động kích hoạt kiểm thử trong GitLab CI / GitHub Actions trên từng PR trước khi merge."
    pf.font.name = FONT_MAIN
    pf.font.size = Pt(9.0)
    pf.font.color.rgb = TEXT_MUTED

    # ================= BLOCK 3 (BOTTOM FULL WIDTH): PLATFORM POSITIONING =================
    b3_card = add_card(x_left, bot_y, bot_w, bot_h, border_color=ACCENT_PURPLE, bg_color=CARD_BG, border_width=1.5, radius=0.05)

    # Header
    b3_tbox = s7.shapes.add_textbox(x_left + Inches(0.35), bot_y + Inches(0.18), Inches(10.0), Inches(0.35))
    tf = b3_tbox.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "🎯 KHI NÀO NÊN CHỌN AIRFLOW HAY DAGSTER?"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    add_badge(x_left + bot_w - Inches(2.30), bot_y + Inches(0.18), Inches(2.00), Inches(0.32),
              "ĐỊNH VỊ NỀN TẢNG", ACCENT_PURPLE, BORDER_PURPLE, CARD_BG_ALT, 10.0)

    # Inner 2 columns
    inner_w = Inches(8.40)
    inner_c1_x = x_left + Inches(0.35)
    inner_c2_x = inner_c1_x + inner_w + Inches(0.54)
    inner_content_y = bot_y + Inches(0.65)

    # Subtle divider line between columns
    divider = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, inner_c1_x + inner_w + Inches(0.26), inner_content_y, Pt(1), Inches(3.90))
    divider.fill.solid()
    divider.fill.fore_color.rgb = BORDER_SUBTLE
    divider.line.fill.background()

    # Column 1: Airflow
    c1_tbox = s7.shapes.add_textbox(inner_c1_x, inner_content_y, inner_w, Inches(0.35))
    tf = c1_tbox.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "⚙️ KHI NÀO NÊN CHỌN APACHE AIRFLOW?"
    p.font.name = FONT_MAIN
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    add_badge(inner_c1_x + inner_w - Inches(1.65), inner_content_y, Inches(1.65), Inches(0.28),
              "Hạ Tầng Sẵn Có", ACCENT_CYAN, BORDER_CYAN, CARD_BG_ALT, 9.5)

    # Bullets Airflow
    c1_bullets_box = s7.shapes.add_textbox(inner_c1_x, inner_content_y + Inches(0.40), inner_w, Inches(2.55))
    tf = c1_bullets_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    
    af_bullets = [
        ("Điều phối hạ tầng chung (Infra Orchestration):", "Bật/tắt máy ảo, trigger job dbt snapshot, gửi email cảnh báo, dọn log hệ thống định kỳ."),
        ("Tác vụ độc lập, thô (Coarse-grained Batch):", "Kích hoạt các job Spark, Trino, Flink chạy độc lập, không cần trao đổi dữ liệu phức tạp."),
        ("Đã có sẵn cụm On-Premise ổn định:", "Doanh nghiệp đã đầu tư hạ tầng VM/Celery với đội ngũ Ops quen vận hành 6 thành phần và NFS."),
        ("Kiểm thử mô hình truyền thống:", "Chấp nhận việc test gắn chặt với Database (dựng Docker Compose mất 3-5 phút/lần test).")
    ]
    for idx, (b_title, b_desc) in enumerate(af_bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {b_title} "
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        if idx > 0:
            p.space_before = Pt(4)
        pr = p.add_run()
        pr.text = b_desc
        pr.font.bold = False
        pr.font.color.rgb = TEXT_MUTED

    # Summary box Airflow
    sum_box_y = inner_content_y + Inches(2.95)
    sum_box_h = Inches(0.85)
    add_card(inner_c1_x, sum_box_y, inner_w, sum_box_h, border_color=BORDER_CYAN, bg_color=CARD_BG_ALT, border_width=1.0, radius=0.04)
    
    sbox_t = s7.shapes.add_textbox(inner_c1_x + Inches(0.18), sum_box_y + Inches(0.12), inner_w - Inches(0.36), sum_box_h - Inches(0.24))
    tf = sbox_t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "📌 Tóm lại: "
    p.font.name = FONT_MAIN
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    pr = p.add_run()
    pr.text = "Phù hợp cho hạ tầng có sẵn, điều phối máy ảo và các luồng xử lý mẻ truyền thống (Cronjob)."
    pr.font.bold = False
    pr.font.color.rgb = TEXT_MAIN

    # Column 2: Dagster
    c2_tbox = s7.shapes.add_textbox(inner_c2_x, inner_content_y, inner_w, Inches(0.35))
    tf = c2_tbox.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "💎 KHI NÀO NÊN CHỌN DAGSTER?"
    p.font.name = FONT_MAIN
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    add_badge(inner_c2_x + inner_w - Inches(1.65), inner_content_y, Inches(1.65), Inches(0.28),
              "ĐỀ XUẤT SỐ 1", ACCENT_GREEN, RGBColor(22, 101, 52), CARD_BG_ALT, 9.5)

    # Bullets Dagster
    c2_bullets_box = s7.shapes.add_textbox(inner_c2_x, inner_content_y + Inches(0.40), inner_w, Inches(2.55))
    tf = c2_bullets_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)

    dg_bullets = [
        ("Thay thế AWS Step Functions On-Premise:", "Pipeline phân nhánh nghiệp vụ phức tạp, xử lý hồ sơ, OCR và đối soát theo State Machine."),
        ("Hạ tầng Cloud-Native K8s tinh gọn:", "Chỉ 3 service cốt lõi (Webserver, Daemon, Postgres nhẹ). Lưu trữ MinIO S3 nội bộ. Cô lập Pods qua gRPC."),
        ("Kiểm thử tự động CI/CD 0.12s:", "Viết unit test in-memory bằng pytest trong RAM, tích hợp mượt mà GitLab CI / GitHub Actions tự động kiểm thử."),
        ("Bảo vệ chất lượng dữ liệu:", "Tích hợp sẵn Asset Checks (blocking=True) chặn đứng hóa đơn vượt ngân sách 100M ngay trong chu trình chạy.")
    ]
    for idx, (b_title, b_desc) in enumerate(dg_bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {b_title} "
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        if idx > 0:
            p.space_before = Pt(4)
        pr = p.add_run()
        pr.text = b_desc
        pr.font.bold = False
        pr.font.color.rgb = TEXT_MUTED

    # Summary box Dagster
    add_card(inner_c2_x, sum_box_y, inner_w, sum_box_h, border_color=RGBColor(22, 101, 52), bg_color=CARD_BG_ALT, border_width=1.0, radius=0.04)

    sbox_dg = s7.shapes.add_textbox(inner_c2_x + Inches(0.18), sum_box_y + Inches(0.12), inner_w - Inches(0.36), sum_box_h - Inches(0.24))
    tf = sbox_dg.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "🚀 Tóm lại: "
    p.font.name = FONT_MAIN
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    pr = p.add_run()
    pr.text = "Nền tảng toàn diện nhất để thay thế AWS Step Functions trên On-Premise, chuẩn hóa văn hóa kỹ nghệ phần mềm cho đội Data!"
    pr.font.bold = False
    pr.font.color.rgb = TEXT_MAIN

    # Speaker notes
    notes_slide = s7.notes_slide
    tf_n = notes_slide.notes_text_frame
    tf_n.text = "Speaker Note: Slide 7 giữ nguyên bố cục 3 khối gốc: (1) Kiến trúc hạ tầng on-premise (Airflow 6 thành phần vs Dagster K8s/MinIO); (2) Kiểm thử CI/CD (Airflow chậm 3-5 phút vs Dagster in-memory 0.12s); (3) Định vị chiến lược khi nào chọn Airflow vs Dagster."

    prs.save(pptx_path)
    print(f"Successfully updated Slide 7 in {pptx_path} with original 3-block layout and standardized design!")

if __name__ == "__main__":
    build_slide7()
